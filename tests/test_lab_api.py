"""End-to-end lab flow through the HTTP API: data -> approved version -> audited run -> journal -> exports -> Pine -> comparison."""
import asyncio
import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from helpers import random_market
from pyapp import jobs
from pyapp.engine.data import iso
from pyapp.main import app

SPEC = {"name": "EMA cross (test)", "timeframe": "M15",
        "series": {"ema": {"type": "ema", "period": 20}, "atr": {"type": "atr", "period": 14}},
        "entry": {"long": {"all": [{"left": "close", "op": ">", "right": "ema"}, {"left": {"s": "close", "shift": 1}, "op": "<=", "right": {"s": "ema", "shift": 1}}]},
                  "short": {"all": [{"left": "close", "op": "<", "right": "ema"}, {"left": {"s": "close", "shift": 1}, "op": ">=", "right": {"s": "ema", "shift": 1}}]}},
        "stop": {"type": "atr", "mult": 1.5, "series": "atr"}, "targets": [{"type": "rr", "value": 1.5}], "risk": {"model": "percent_equity", "value": 1.0}}


def csv_bytes(n=1200, tf="M15", seed=7):
    cs = random_market(n, tf, seed)
    lines = ["time,open,high,low,close,volume"]
    lines += [f"{iso(cs.t[i])},{cs.o[i]},{cs.h[i]},{cs.l[i]},{cs.c[i]},100" for i in range(len(cs))]
    return "\n".join(lines).encode()


def run_all_jobs():
    async def go():
        for job in reversed(await jobs.db.list("jobs")):
            if job["status"] == "QUEUED":
                await jobs.run_job(job)
    asyncio.run(go())


def setup_run(client):
    data = client.post("/api/market-data/csv", files={"file": ("eurusd.csv", csv_bytes(), "text/csv")}, data={"symbol": "EURUSD", "timeframe": "M15", "provider": "Test vendor", "tz": "UTC"})
    assert data.status_code == 201, data.text
    strategy = client.post("/api/library/custom", json={"name": "EMA cross"}).json()
    version = client.post("/api/versions", json={"strategyId": strategy["id"], "spec": SPEC}).json()
    assert version["status"] == "DRAFT" and version["readiness"]["fidelity"] == "USER_AUTHORED_RULES"
    return data.json(), strategy, version


def test_end_to_end_flow_with_exports_and_reconciliation():
    client = TestClient(app)
    dataset, strategy, version = setup_run(client)
    config = {"dataset_id": dataset["id"], "symbol": "EURUSD", "initial_capital": 10000, "costs": {"spread": 0.0001, "slippage": 0.00001, "commission_model": "per_lot_side", "commission_value": 3.5},
              "causality_samples": 3, "start": "2024-01-03"}
    assert client.post("/api/runs", json={"versionId": version["id"], "config": config}).status_code == 409           # draft versions cannot be backtested
    approved = client.post(f"/api/versions/{version['id']}/approve", json={"acknowledged": []})
    assert approved.status_code == 200 and approved.json()["status"] == "APPROVED"
    assert client.put(f"/api/versions/{version['id']}", json={"spec": SPEC}).status_code == 409                       # approved versions are immutable
    queued = client.post("/api/runs", json={"versionId": version["id"], "config": config, "label": "baseline"})
    assert queued.status_code == 202
    run_all_jobs()
    job = client.get(f"/api/jobs/{queued.json()['id']}").json()
    assert job["status"] == "COMPLETED", job
    run_id = job["result"]["runId"]
    run = client.get(f"/api/runs/{run_id}").json()
    assert run["status"] == "VALID", [c for c in run["validation"]["checks"] if c["status"] == "FAIL"]
    assert run["strategyVersionId"] == version["id"] and run["fidelity"] == "USER_AUTHORED_RULES"
    assert run["lookaheadAudit"]["passed"] and run["reproducibility"]["identical"]
    assert client.get(f"/api/versions/{version['id']}").json()["status"] == "FROZEN"                                    # frozen on first use
    assert run["metrics"]["totalTrades"] > 5

    listing = client.get(f"/api/runs/{run_id}/trades?limit=1000").json()
    assert listing["total"] == run["tradeCount"]
    first = listing["trades"][0]
    detail = client.get(f"/api/runs/{run_id}/trades/{first['id']}").json()
    assert detail["explanation"] and "Stop-loss" in " ".join(detail["explanation"])
    chart = client.get(f"/api/runs/{run_id}/trades/{first['id']}/chart.svg")
    assert chart.status_code == 200 and chart.text.startswith("<svg") and "Entry" in chart.text

    csv_response = client.get(f"/api/runs/{run_id}/export.csv")
    assert csv_response.status_code == 200
    csv_lines = csv_response.content.decode("utf-8-sig").strip().splitlines()
    assert len(csv_lines) == run["tradeCount"] + 1
    xlsx = client.get(f"/api/runs/{run_id}/export.xlsx")
    workbook = load_workbook(io.BytesIO(xlsx.content))
    assert workbook.sheetnames == ["Trade Journal", "Performance Summary", "Strategy Rules", "Backtest Configuration", "Validation & Exceptions"]
    journal = workbook["Trade Journal"]
    header = [c.value for c in journal[1]]
    net_column = header.index("Net P&L") + 1
    ids = [journal.cell(r, 1).value for r in range(2, journal.max_row + 1)]
    assert len(ids) == len(set(ids)) == run["tradeCount"]                                                              # every trade exactly once
    total = sum(journal.cell(r, net_column).value for r in range(2, journal.max_row + 1))
    assert abs(total - (run["finalBalance"] - run["initialCapital"])) < 0.05 * run["tradeCount"] + 0.05                 # row rounding only

    pine = client.post(f"/api/runs/{run_id}/pine").json()
    assert pine["script"].startswith("//@version=6") and pine["audit"]["verdict"] in ("PASS", "WARN")
    assert client.post(f"/api/versions/{version['id']}/pine", json={}).json()["versionId"] == version["id"]

    second = client.post("/api/runs", json={"versionId": version["id"], "config": {**config, "costs": {"spread": 0.0003, "slippage": 0.00003}}, "label": "wider costs"})
    run_all_jobs()
    second_id = client.get(f"/api/jobs/{second.json()['id']}").json()["result"]["runId"]
    compared = client.post("/api/compare", json={"runIds": [run_id, second_id]}).json()
    nets = {r["runId"]: r["netPnl"] for r in compared["runs"]}
    assert nets[second_id] < nets[run_id]                                                                               # higher costs never help


def test_invalid_runs_cannot_be_exported_as_valid():
    client = TestClient(app)
    dataset, strategy, version = setup_run(client)
    client.post(f"/api/versions/{version['id']}/approve", json={"acknowledged": []})
    queued = client.post("/api/runs", json={"versionId": version["id"], "config": {"dataset_id": dataset["id"], "symbol": "EURUSD", "costs": {"spread": 0.0001}}})
    run_all_jobs()
    run_id = client.get(f"/api/jobs/{queued.json()['id']}").json()["result"]["runId"]
    async def tamper():
        from pyapp.db import db
        run = await db.get("bt_runs", run_id)
        run["status"] = "INVALID"
        run["validation"]["checks"].append({"name": "Injected failure", "status": "FAIL", "detail": ""})
        await db.save("bt_runs", run)
    asyncio.run(tamper())
    assert client.get(f"/api/runs/{run_id}/export.csv").status_code == 409
    assert client.get(f"/api/runs/{run_id}/export.xlsx").status_code == 409


def test_assumptions_must_be_acknowledged_before_approval():
    client = TestClient(app)
    _, strategy, _ = setup_run(client)
    version = client.post("/api/versions", json={"strategyId": strategy["id"], "spec": SPEC, "assumptions": [{"component": "stop", "text": "ATR x1.5 chosen by me", "reason": "needed"}]}).json()
    assumption_id = version["assumptions"][0]["id"]
    assert client.post(f"/api/versions/{version['id']}/approve", json={"acknowledged": []}).status_code == 409
    approved = client.post(f"/api/versions/{version['id']}/approve", json={"acknowledged": [assumption_id]})
    assert approved.status_code == 200 and approved.json()["assumptions"][0]["approved"]


def test_spec_validation_reports_readable_errors():
    client = TestClient(app)
    bad = {**SPEC, "entry": {"long": {"left": "close", "op": ">", "right": "no_such_series"}}}
    result = client.post("/api/spec/validate", json={"spec": bad}).json()
    assert not result["valid"] and "no_such_series" in result["errors"][0]
    assert client.post("/api/spec/validate", json={"spec": SPEC}).json()["valid"]
    for name, template in client.get("/api/spec/templates").json().items():
        assert client.post("/api/spec/validate", json={"spec": template}).json()["valid"], name


def test_optimization_creates_separately_labelled_versions_and_reports_oos():
    client = TestClient(app)
    dataset, strategy, version = setup_run(client)
    client.post(f"/api/versions/{version['id']}/approve", json={"acknowledged": []})
    config = {"dataset_id": dataset["id"], "symbol": "EURUSD", "costs": {"spread": 0.0001}, "start": "2024-01-03"}
    job = client.post("/api/optimize", json={"versionId": version["id"], "config": config, "grid": {"stop.mult": [1.0, 2.0], "targets.0.value": [1.0, 2.0]}, "minTrades": 3, "topK": 2})
    assert job.status_code == 202
    run_all_jobs()
    finished = client.get(f"/api/jobs/{job.json()['id']}").json()
    assert finished["status"] == "COMPLETED", finished
    result = client.get(f"/api/optimizations/{finished['result']['optimizationId']}").json()
    assert len(result["trials"]) == 4 and result["top"] and "outOfSample" in result["top"][0]
    assert "over-fit" in result["warning"].lower() or "biased" in result["warning"].lower()
    variants = [v for v in client.get(f"/api/versions?strategyId={strategy['id']}").json() if v["kind"] == "optimized_variant"]
    assert variants and all(v["label"].startswith("opt-") and v["status"] == "DRAFT" and v["parentVersionId"] == version["id"] for v in variants)
    assert client.get(f"/api/versions/{version['id']}").json()["spec"]["stop"]["mult"] == 1.5                           # original untouched
    assert client.post("/api/optimize", json={"versionId": version["id"], "config": config, "grid": {"stop.mult": list(range(1, 30)), "targets.0.value": list(range(1, 10))}}).status_code == 422

    walk = client.post("/api/optimize", json={"versionId": version["id"], "config": config, "grid": {"stop.mult": [1.0, 2.0]}, "mode": "walk_forward", "folds": 2, "minTrades": 1})
    run_all_jobs()
    done = client.get(f"/api/jobs/{walk.json()['id']}").json()
    assert done["status"] == "COMPLETED", done
    folds = client.get(f"/api/optimizations/{done['result']['optimizationId']}").json()["folds"]
    assert len(folds) == 2 and all(f["train"][1] <= f["test"][0] for f in folds)                                         # tests always follow training


def test_csv_ingestion_is_strict_and_discloses_repairs():
    client = TestClient(app)
    good = client.post("/api/market-data/csv", files={"file": ("a.csv", csv_bytes(300), "text/csv")}, data={"symbol": "EURUSD", "timeframe": "M15", "provider": "xx"})
    assert good.status_code == 201 and good.json()["quality"]["coveragePct"] > 90
    broken = csv_bytes(300).replace(b"1.1", b"1.1", 1) + b"\n2024-01-01T00:00:00Z,1,2,3,4,5"
    assert client.post("/api/market-data/csv", files={"file": ("b.csv", broken, "text/csv")}, data={"symbol": "EURUSD", "timeframe": "M15", "provider": "xx"}).status_code == 400
    wrong_tf = client.post("/api/market-data/csv", files={"file": ("c.csv", csv_bytes(300, "M5"), "text/csv")}, data={"symbol": "EURUSD", "timeframe": "H4", "provider": "xx"})
    assert wrong_tf.status_code == 400 and "aligned" in wrong_tf.json()["detail"]                                         # M5 stamps are not H4 boundaries
